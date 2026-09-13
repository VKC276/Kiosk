#!/usr/bin/env python3
"""Importera 10-kort från Excel (Kortnummer, Antal kvarvarande besök, Status, Senast klippt)."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parents[1]


def load_config() -> dict:
    path = ROOT / "config.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def rows_from_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("Installera openpyxl: pip install openpyxl") from exc

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not values or values[0] is None or str(values[0]).strip() == "":
            continue
        item: dict = {}
        for key, value in zip(headers, values):
            if not key:
                continue
            if hasattr(value, "strftime"):
                item[key] = value.strftime("%Y-%m-%d %H:%M:%S")
            else:
                item[key] = value
        rows.append(item)
    return rows


def main() -> int:
    cfg = load_config()
    cf = cfg.get("CLOUDFLARE") or {}
    parser = argparse.ArgumentParser(description="Importera 10-kort.xlsx till Cloudflare")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--api-url", default=cf.get("apiUrl") or "")
    parser.add_argument("--admin-token", default=cf.get("adminToken") or "")
    args = parser.parse_args()

    if not args.api_url or not args.admin_token:
        print("Saknar --api-url och/eller --admin-token.", file=sys.stderr)
        return 1
    if not args.xlsx.is_file():
        print(f"Hittar inte {args.xlsx}", file=sys.stderr)
        return 1

    rows = rows_from_xlsx(args.xlsx)
    print(f"Läste {len(rows)} 10-kort från {args.xlsx}")
    response = requests.post(
        f"{str(args.api_url).rstrip('/')}/api/admin/import/tencards",
        json=rows,
        headers={"Authorization": f"Bearer {args.admin_token}", "Content-Type": "application/json"},
        timeout=60,
    )
    if not response.ok:
        print(f"Importfel: HTTP {response.status_code} {response.text[:300]}", file=sys.stderr)
        return 1
    print(response.json())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
